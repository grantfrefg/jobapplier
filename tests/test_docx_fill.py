"""Your own Word resume, filled with the tailored content (Resume look -> My Word file)."""

from pathlib import Path

import pytest

from applypilot import config, own_docs
from applypilot.scoring import docx_fill, look, pdf

FIXTURE = Path(__file__).parent / "fixtures" / "word_resume.docx"


def paragraphs(path):
    from docx import Document
    return [(p.style.name, [(r.text, r.bold) for r in p.runs]) for p in Document(str(path)).paragraphs]


def test_tailored_content_goes_into_your_own_formatting(tmp_path):
    out = tmp_path / "out.docx"
    assert docx_fill.fill(look.SAMPLE, out, FIXTURE) == {"summary": True, "skills": True, "experience": 2}
    text = docx_fill.text_of(out)
    assert "Data scientist with 7 years" in text and "Old summary" not in text
    assert "Old bullet" not in text and "Technical lead for production recommender systems" in text
    assert "Kept bullet from the bank" in text  # an entry tailoring left out keeps your bullets
    assert "Jan 2021 – Present" in text and "Master of Data Science" in text  # facts as you wrote them
    assert "Working holiday visa" not in text  # automatic: a limited visa is left off the header
    rows = paragraphs(out)
    assert ("Normal", [("Machine learning: ", True), ("PyTorch, scikit-learn, recommender systems, A/B testing",
                                                      None)]) in rows  # your bold category style
    bullets = [r for r in rows if r[0] == "List Bullet"]
    assert len(bullets) == 4 and all(len(runs) == 1 for _, runs in bullets)


def test_a_document_without_sections_or_not_word_is_refused(tmp_path):
    from docx import Document
    plain = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("Just a letter with no headings")
    doc.save(str(plain))
    with pytest.raises(docx_fill.FillError):
        docx_fill.fill(look.SAMPLE, tmp_path / "out.docx", plain)
    with pytest.raises(docx_fill.FillError):
        docx_fill.check_docx(b"%PDF-1.4 not word")


def test_the_word_copy_is_made_and_sent_when_chosen(tmp_path, monkeypatch):
    monkeypatch.setattr(pdf, "render_pdf", lambda html, out: Path(out).write_bytes(b"%PDF-1.4"))
    config.APP_DIR.mkdir(parents=True, exist_ok=True)
    docx_fill.template_path().write_bytes(FIXTURE.read_bytes())
    resume = tmp_path / "Acme_Data.txt"
    resume.write_text(look.SAMPLE, encoding="utf-8")
    job = {"tailored_resume_path": str(resume)}

    pdf.convert_to_pdf(resume)  # another layout: PDF only
    assert not resume.with_suffix(".docx").exists() and own_docs.documents(job)["resume"]["file"].suffix == ".pdf"

    monkeypatch.setenv("RESUME_LAYOUT", "word")
    pdf.convert_to_pdf(resume)
    assert resume.with_suffix(".pdf").exists() and resume.with_suffix(".docx").exists()
    doc = own_docs.documents(job)["resume"]
    assert doc["file"].suffix == ".docx" and "Technical lead" in doc["text"]
